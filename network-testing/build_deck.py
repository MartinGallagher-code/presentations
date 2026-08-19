#!/usr/bin/env python3
"""Build network_testing.pptx from network_testing_talk.md using python-pptx.

Deliberately plain: default template layouts, title + bullets + tables +
monospace code blocks, full speaker notes. Maximum PowerPoint compatibility.
"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

MD = 'network_testing_talk.md'
OUT = 'network_testing.pptx'

INK = RGBColor(0x1C, 0x2B, 0x3A)
ACCENT = RGBColor(0x3E, 0x5C, 0x76)
CODE = RGBColor(0x24, 0x35, 0x4A)

def clean(s):
    s = s.replace('**', '').replace('`', '')
    return s.strip()

# ---- parse markdown into slides -------------------------------------------
slides = []   # each: dict(kind, title, blocks, notes)
raw = open(MD).read().split('\n---\n')

for chunk in raw:
    lines = [l.rstrip() for l in chunk.strip('\n').split('\n')]
    if not lines or not any(l.strip() for l in lines):
        continue
    first = next(l for l in lines if l.strip())
    if first.startswith('# '):           # title slide
        title = clean(first[2:])
        sub = ''
        for l in lines[1:]:
            if l.startswith('### '):
                sub = clean(l[4:])
                break
        slides.append(dict(kind='title', title=title, sub=sub, blocks=[], notes=''))
        continue
    if first.startswith('## '):          # section divider
        slides.append(dict(kind='section', title=clean(first[3:]), blocks=[], notes=''))
        continue
    if not first.startswith('### Slide:'):
        continue
    title = clean(first[len('### Slide:'):])
    blocks, notes = [], ''
    i = lines.index(first) + 1
    in_code, code_lines = False, []
    table_rows = []
    def flush_table():
        if table_rows:
            blocks.append(('table', list(table_rows)))
            table_rows.clear()
    while i < len(lines):
        l = lines[i]
        if l.strip().startswith('```'):
            flush_table()
            if in_code:
                blocks.append(('code', code_lines))
                code_lines = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_lines.append(l)
            i += 1
            continue
        if l.strip().startswith('**Notes:**'):
            notes = clean(l.strip()[len('**Notes:**'):])
            i += 1
            while i < len(lines) and lines[i].strip():
                notes += ' ' + clean(lines[i])
                i += 1
            continue
        if l.strip().startswith('|'):
            cells = [clean(c) for c in l.strip().strip('|').split('|')]
            if not all(re.fullmatch(r'-{3,}', c) for c in cells):
                table_rows.append(cells)
            i += 1
            continue
        if l.strip().startswith('=> '):
            flush_table()
            items = [clean(seg) for seg in l.strip()[3:].split(' -> ')]
            blocks.append(('flow', items))
            i += 1
            continue
        flush_table()
        m = re.match(r'^(\s*)- (.*)$', l)
        if m:
            lvl = 1 if len(m.group(1)) >= 2 else 0
            blocks.append(('bullet', lvl, clean(m.group(2))))
        elif l.strip():
            blocks.append(('para', clean(l)))
        i += 1
    flush_table()
    if in_code and code_lines:
        blocks.append(('code', code_lines))
    slides.append(dict(kind='content', title=title, blocks=blocks, notes=notes))

# ---- build the deck --------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
L_TITLE, L_CONTENT, L_SECTION, L_TITLE_ONLY = 0, 1, 2, 5

def set_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text

for sl in slides:
    if sl['kind'] == 'title':
        s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
        s.shapes.title.text = sl['title']
        tp = s.shapes.title.text_frame.paragraphs[0]
        tp.font.size = Pt(44)
        tp.font.bold = True
        tp.font.color.rgb = INK
        stf = s.placeholders[1].text_frame
        p1 = stf.paragraphs[0]
        p1.text = sl['sub']
        p1.font.size = Pt(20)
        p1.font.color.rgb = ACCENT
        p2 = stf.add_paragraph()
        p2.text = 'and reachable, keeping the server list honest'
        p2.font.size = Pt(15)
        p2.font.italic = True
        p3 = stf.add_paragraph()
        p3.text = ''
        p4 = stf.add_paragraph()
        p4.text = 'Martin Gallagher'
        p4.font.size = Pt(14)
        p4.font.bold = True
        set_notes(s, sl['notes'])
        continue
    if sl['kind'] == 'section':
        s = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
        s.shapes.title.text = sl['title']
        set_notes(s, sl['notes'])
        continue

    has_table = any(b[0] == 'table' for b in sl['blocks'])
    has_code = any(b[0] == 'code' for b in sl['blocks'])
    has_flow = any(b[0] == 'flow' for b in sl['blocks'])
    layout = L_TITLE_ONLY if (has_table or has_code or has_flow) else L_CONTENT
    s = prs.slides.add_slide(prs.slide_layouts[layout])
    s.shapes.title.text = sl['title']
    t = s.shapes.title.text_frame.paragraphs[0]
    t.font.size = Pt(28)
    t.font.bold = True
    t.font.color.rgb = INK

    y = Inches(1.35)
    if layout == L_CONTENT:
        body = s.placeholders[1]
        tf = body.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        first = True
        for b in sl['blocks']:
            if b[0] == 'bullet':
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.text = b[2]
                p.level = b[1]
                p.font.size = Pt(16 if b[1] == 0 else 14)
            elif b[0] == 'para':
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.text = b[1]
                p.font.size = Pt(16)
        set_notes(s, sl['notes'])
        continue

    # mixed slide: place code / table / bullets manually
    for b in sl['blocks']:
        if b[0] == 'code':
            n = len(b[1])
            h = Inches(0.25 + 0.24 * n)
            box = s.shapes.add_textbox(Inches(0.7), y, Inches(12.0), h)
            tf = box.text_frame
            tf.word_wrap = False
            for j, line in enumerate(b[1]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = line if line else ' '
                p.font.name = 'Consolas'
                p.font.size = Pt(13)
                p.font.color.rgb = CODE
            y = y + h + Inches(0.15)
        elif b[0] == 'table':
            rows, cols = len(b[1]), len(b[1][0])
            h = Inches(0.4 * rows)
            shp = s.shapes.add_table(rows, cols, Inches(0.7), y, Inches(12.0), h)
            tbl = shp.table
            for ri, row in enumerate(b[1]):
                for ci in range(cols):
                    cell = tbl.cell(ri, ci)
                    cell.text = row[ci] if ci < len(row) else ''
                    for p in cell.text_frame.paragraphs:
                        p.font.size = Pt(12 if ri else 12.5)
                        p.font.bold = (ri == 0)
            y = y + h + Inches(0.2)
        elif b[0] == 'flow':
            items = b[1]
            n = len(items)
            gap = 0.42
            bw = (12.0 - gap * (n - 1)) / n
            bh = Inches(1.0)
            colors = [RGBColor(0x3D, 0x8B, 0x4F), RGBColor(0x12, 0x89, 0x7B),
                      RGBColor(0xD9, 0x7C, 0x1F), RGBColor(0xBE, 0x3B, 0x4E),
                      RGBColor(0x12, 0x89, 0x7B), RGBColor(0x3E, 0x5C, 0x76)]
            for k, item in enumerate(items):
                parts = item.split('|')
                cmd = parts[0].strip()
                cap = parts[1].strip() if len(parts) > 1 else ''
                x = Inches(0.7 + k * (bw + gap))
                box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(bw), bh)
                box.fill.solid()
                box.fill.fore_color.rgb = colors[k % len(colors)]
                box.line.fill.background()
                btf = box.text_frame
                btf.word_wrap = True
                bp = btf.paragraphs[0]
                bp.text = cmd
                bp.font.size = Pt(11.5)
                bp.font.bold = True
                bp.font.name = 'Consolas'
                bp.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if cap:
                    cp = btf.add_paragraph()
                    cp.text = cap
                    cp.font.size = Pt(10)
                    cp.font.color.rgb = RGBColor(0xF0, 0xF3, 0xF6)
                if k < n - 1:
                    ar = s.shapes.add_textbox(Inches(0.7 + (k + 1) * bw + k * gap), y + Inches(0.28), Inches(gap), Inches(0.4))
                    ap = ar.text_frame.paragraphs[0]
                    ap.text = '\u2192'
                    ap.font.size = Pt(16)
                    ap.font.bold = True
                    ap.font.color.rgb = ACCENT
            y = y + bh + Inches(0.22)
        elif b[0] in ('bullet', 'para'):
            txt = b[2] if b[0] == 'bullet' else b[1]
            lvl = b[1] if b[0] == 'bullet' else 0
            est = max(1, (len(txt) // 110) + 1)
            h = Inches(0.34 * est)
            box = s.shapes.add_textbox(Inches(0.9 + 0.35 * lvl), y, Inches(11.6 - 0.35 * lvl), h)
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ('• ' if b[0] == 'bullet' else '') + txt
            p.font.size = Pt(14 if lvl == 0 else 12.5)
            y = y + h + Inches(0.08)
    set_notes(s, sl['notes'])

prs.save(OUT)
print('wrote %s with %d slides' % (OUT, len(prs.slides.slides if hasattr(prs.slides,"slides") else prs.slides._sldIdLst)))
